# src/visualization/report_builder.py
"""
Report builder: produces PDF and PowerPoint reports containing:
 - Title and summary text
 - Charts (PNG)
 - Formatted tables (from pandas DataFrames)
 - Metrics dictionary (key: value)
 - Placeholder for AI-written summary (integrate phi-3b later)

Dependencies:
  pip install reportlab python-pptx pandas matplotlib
"""

from pathlib import Path
from typing import List, Dict, Any
import io

from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table as RLTable, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors

from pptx import Presentation
from pptx.util import Inches, Pt

import pandas as pd


# -------------------------
# Helpers
# -------------------------
def _df_to_rl_table_data(df: pd.DataFrame, include_index: bool = False) -> List[List[Any]]:
    """
    Convert pandas DataFrame to data array consumable by reportlab Table.
    """
    if include_index:
        headers = [str(df.index.name or "")] + [str(c) for c in df.columns]
    else:
        headers = [str(c) for c in df.columns]

    data = [headers]

    for i, row in df.iterrows():
        cells = []
        if include_index:
            cells.append(str(i))
        for c in df.columns:
            val = row[c]
            # Convert floats/numbers to readable strings
            if pd.isna(val):
                cells.append("")
            else:
                cells.append(str(val))
        data.append(cells)
    return data


# -------------------------
# AI summary placeholder
# -------------------------
def generate_ai_summary(metrics: Dict[str, Any], top_n: int = 5) -> str:
    """
    Placeholder for AI-written summary. For now, creates a deterministic short summary
    from metrics. Later this should call phi-3b/RAG to produce richer text.
    """
    lines = []
    lines.append("Executive summary:")
    if not metrics:
        lines.append("No metrics available.")
    else:
        for k, v in list(metrics.items())[:top_n]:
            try:
                if isinstance(v, float):
                    lines.append(f"- {k.replace('_', ' ').title()}: {v:.3f}")
                else:
                    lines.append(f"- {k.replace('_', ' ').title()}: {v}")
            except Exception:
                lines.append(f"- {k}: {v}")
    return "\n".join(lines)


# -------------------------
# PDF Report Builder
# -------------------------
def build_pdf_report(
    out_path: str | Path,
    title: str,
    cleaned_tables: Dict[str, pd.DataFrame],
    metrics: Dict[str, Any],
    chart_paths: List[str] | None = None,
    author: str | None = None,
):
    """
    Build a simple PDF report including:
     - Title
     - AI summary (from metrics)
     - Charts (PNG files)
     - Tables (pandas DataFrames)
     - Metrics list
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(str(out_path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = styles["Title"]
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 12))

    # Author
    if author:
        author_style = ParagraphStyle(name="Author", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
        story.append(Paragraph(f"Author: {author}", author_style))
        story.append(Spacer(1, 8))

    # AI summary
    summary_text = generate_ai_summary(metrics)
    story.append(Paragraph(summary_text.replace("\n", "<br/>"), styles["Normal"]))
    story.append(Spacer(1, 12))

    # Metrics key-values
    if metrics:
        kv_data = [["Metric", "Value"]]
        for k, v in metrics.items():
            kv_data.append([str(k), str(v)])
        tbl = RLTable(kv_data, hAlign="LEFT", colWidths=[200, 300])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 12))

    # Charts
    if chart_paths:
        for p in chart_paths:
            p = Path(p)
            if p.exists():
                # constrain width to page width - margins
                img = RLImage(str(p), width=450, height=250)
                story.append(img)
                story.append(Spacer(1, 12))

    # Tables
    for name, df in cleaned_tables.items():
        story.append(Paragraph(f"Table: {name}", styles["Heading3"]))
        # convert df to table data
        tbl_data = _df_to_rl_table_data(df)
        # limit size: if table too large, show top 10 rows
        if len(tbl_data) > 12:
            tbl_data = [tbl_data[0]] + tbl_data[1:11]  # headers + first 10 rows
            tbl_data.append(["...", "..."] * (len(tbl_data[0]) // 2))

        tbl = RLTable(tbl_data, hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 12))

    doc.build(story)
    return str(out_path)


# -------------------------
# PPTX Report Builder
# -------------------------
def build_pptx_report(
    out_path: str | Path,
    title: str,
    cleaned_tables: Dict[str, pd.DataFrame],
    metrics: Dict[str, Any],
    chart_paths: List[str] | None = None,
    author: str | None = None,
):
    """
    Build a PowerPoint deck with:
     - Title slide
     - Metrics slide
     - One chart per slide
     - Table slides (first N rows)
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    # Title slide layout
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title
    if author:
        subtitle = slide.placeholders[1]
        subtitle.text = f"Author: {author}"

    # Metrics slide
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # blank-ish
    left = Inches(0.5)
    top = Inches(0.5)
    width = Inches(9)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = "Key Metrics"
    for k, v in metrics.items():
        p = tf.add_paragraph()
        p.text = f"{k}: {v}"
        p.level = 1

    # Charts
    if chart_paths:
        for chart in chart_paths:
            p = Path(chart)
            if p.exists():
                slide = prs.slides.add_slide(prs.slide_layouts[6])  # picture layout / blank
                left = Inches(1)
                top = Inches(1)
                slide.shapes.add_picture(str(p), left, top, width=Inches(8))

    # Tables: add a slide per table (first 10 rows)
    for name, df in cleaned_tables.items():
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        title_shape = slide.shapes.title if slide.shapes.title else None
        if title_shape:
            title_shape.text = f"Table: {name}"
        rows = min(len(df) + 1, 11)  # header + up to 10 rows
        cols = len(df.columns)
        if cols == 0:
            continue
        left = Inches(0.5)
        top = Inches(1)
        width = Inches(9)
        height = Inches(4.5)
        table = slide.shapes.add_table(rows, cols, left, top, width, height).table
        # header
        for j, col in enumerate(df.columns):
            table.cell(0, j).text = str(col)
        for i, (_, row) in enumerate(df.head(10).iterrows(), start=1):
            for j, col in enumerate(df.columns):
                table.cell(i, j).text = str(row[col])

    prs.save(str(out_path))
    return str(out_path)